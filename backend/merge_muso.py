# Merge a Muso export into songs, keeping every detail.
# Output sheets:
#   Summary         - counts and the songs with the most versions
#   Merged by Song  - one row per song (composition). Each column holds the
#                     primary version's value; other versions' distinct values
#                     are comma-joined in "(Duplicates)"
#   Song Versions   - every source row, untouched, grouped under its Song ID
#   Merged by ISRC  - one row per recording (ISRC), same column logic
#   Muso Credits    - the source sheet as-is
import re
import sys
import difflib
import unicodedata
import openpyxl
from collections import OrderedDict, Counter, defaultdict
from openpyxl.styles import Font, PatternFill
from merge_isrc import format_plain, format_merged, HDR_FILL, HDR_FONT, BORDER

MUSO_SHEET = "Muso Credits"
TITLE = "Name of the Song / Title"
PLACEHOLDERS = {"", "N/A", "NA", "NONE", "-", "UNAVAILABLE", "NOT PROVIDED BY MUSO"}
# Columns holding comma-separated names: merged name by name, not as whole strings
LIST_COLS = {
    "Main Artist (All)", "Singer", "Lyricist", "Composer", "Producer",
    "Recording Engineer", "Mixing Engineer", "Mastering Engineer", "Featured Artist",
    "Publisher", "Record Label (credited org)", "Other Credits", "Artist's Role(s)", "Genre", "Author",
}
ISRC_RE = re.compile(r"[A-Z]{2}[A-Z0-9]{3}\d{7}")

# A credit on more than this share of all recordings (the report's own artist)
# says nothing about two recordings being the same song, so it is ignored.
COMMON_CREDIT_SHARE = 0.2
# Title similarity needed to join two different spellings (with a shared credit)
FUZZY_RATIO = 0.89

# Words marking a title part as a version tag rather than part of the song name
VERSION_WORDS = [
    "remix", "re-mix", "mix", "version", "from", "lofi", "lo-fi", "lo fi", "flip",
    "jhankar", "jhankaar", "jhankar beats", "beats", "hifi", "stereo", "reprise", "live",
    "unplugged", "acoustic", "instrumental", "karaoke", "slowed", "reverb", "sped up",
    "revisited", "revisit", "recreated", "recreation", "reloaded", "remastered",
    "remaster", "edit", "cover", "male", "female", "sad", "happy", "duet", "solo", "club",
    "trap", "house", "phonk", "bolly", "afro", "drill", "edm", "dj", "rmx", "vibe",
    "vibes", "trending", "unwind", "soundtrack", "original motion picture", "ost",
    "film", "theme", "guitar", "piano", "flute", "violin", "sax", "saxophone", "chill",
    "dance", "rework", "rewind", "2.0", "3.0", "dialogue", "part", "pt",
]
VERSION_RE = re.compile(r"(?<![a-z])(" + "|".join(re.escape(w) for w in VERSION_WORDS) + r")(?![a-z])", re.I)
TRAILING_RE = re.compile(r"(\s+(" + VERSION_RE.pattern + r")(\s+\S+)?)\s*$", re.I)
# Medleys/mashups hold several songs: never folded into one song by title
MEDLEY_RE = re.compile(r"medley|mash-?up|smashup|jukebox|non[\s-]?stop|commentry|commentary|\sx\s|\bvs\b", re.I)

VERSION_TYPES = [
    ("Medley / Mashup", MEDLEY_RE),
    ("Lofi", re.compile(r"lo[\s-]?fi", re.I)),
    ("Jhankar", re.compile(r"jh?ank+aa?r", re.I)),
    ("Slowed / Reverb", re.compile(r"slowed|reverb|sped up", re.I)),
    ("Remix", re.compile(r"remix|rmx|\bmix\b|club|trap|house|phonk|afro|drill|edm|\bdj\b|rework", re.I)),
    ("Reprise", re.compile(r"reprise", re.I)),
    ("Unplugged / Acoustic", re.compile(r"unplugged|acoustic", re.I)),
    ("Live", re.compile(r"\blive\b|concert", re.I)),
    ("Instrumental", re.compile(r"instrumental|karaoke|guitar|piano|flute|violin|sax|theme", re.I)),
    ("Recreated / Cover", re.compile(r"recreat|revisit|reloaded|cover|tribute|2\.0|3\.0", re.I)),
    ("Male / Female Version", re.compile(r"female|male", re.I)),
    ("Sad / Happy Version", re.compile(r"\bsad\b|\bhappy\b", re.I)),
]

# Common Hinglish spelling variants, per word
WORD_FIX = {"hain": "hai", "nahin": "nahi", "nahi": "nahi", "mein": "me", "main": "me",
            "mai": "me", "hoon": "hu", "hun": "hu", "kyun": "kyu", "kyon": "kyu", "tum": "tum"}


def is_placeholder(v):
    return v is None or str(v).strip().upper() in PLACEHOLDERS


def ascii_lower(s):
    return unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode().lower()


def phonetic(s):
    """Spelling-insensitive key: 'Hungamaa Ho Gaya' == 'Hungama Hogaya',
    'Ek Pyar Ka Naghma' == 'Ek Pyaar Ka Nagma'."""
    words = re.sub(r"[^a-z0-9]+", " ", ascii_lower(s)).split()
    s = "".join(WORD_FIX.get(w, w) for w in words)
    s = re.sub(r"(.)\1+", r"\1", s)  # "Peechhe" -> "Peche" before dropping h's
    s = s.replace("ph", "f").replace("ck", "k")
    s = re.sub(r"([bcdgjkpstr])h", r"\1", s)
    s = s.replace("w", "v").replace("z", "j").replace("q", "k")
    s = s.replace("ee", "i").replace("oo", "u").replace("y", "i")
    return re.sub(r"(.)\1+", r"\1", s)


def name_key(s):
    """Same name written differently (case, dashes, dots, spacing)."""
    return re.sub(r"[^a-z0-9]", "", ascii_lower(s))


def is_medley(title):
    return bool(MEDLEY_RE.search(f" {title} "))


def clean_title(title):
    t = str(title or "")
    t = re.sub(r"\([^)]*\)|\[[^\]]*\]", " ", t)  # (From "Queen"), [Remix]
    t = re.sub(r",?\s*\b(pt|part)\.?\s*\d+\b", " ", t, flags=re.I)  # ", Pt. 1"
    m = re.search(r"\s+-\s+[^-/]*/(.*)$", t)  # "Song - Film / Soundtrack Version"
    if m and VERSION_RE.search(m.group(1)):
        t = t[:m.start()]
    parts = re.split(r"\s+[-–—|:]\s+|\s*/\s*|\s+-(?=\S)|(?<=\S)-\s+", t)
    # Cut at the first part that is a version tag: "Song - Film / Soundtrack Version",
    # "Song - From Dil Hai Betaab - Remix", "Song - Film- Piya Ka Ghar"
    for n in range(1, len(parts)):
        if VERSION_RE.search(parts[n]) or "/" in parts[n]:
            parts = parts[:n]
            break
    t = " ".join(parts)
    prev = None
    while prev != t:  # "Tera Naam Liya LoFi Flip", "Ek Hasina Thi Lofi Mix"
        prev = t
        m = TRAILING_RE.search(t)
        if m and m.start() > 0:
            t = t[:m.start()]
    return t.strip(" -–—|:/,.")


def title_key(title):
    if is_medley(title):
        return None
    return phonetic(clean_title(title)) or None


def version_type(title):
    t = str(title or "")
    for label, rx in VERSION_TYPES:
        if rx.search(t):
            return label
    return "Original / Re-release"


def isrc_tokens(v):
    return ISRC_RE.findall(str(v or "").upper())


def split_names(v):
    if is_placeholder(v):
        return []
    return [n.strip() for n in str(v).split(",") if not is_placeholder(n)]


CREDIT_JUNK = {"instrumental", "unknown", "various", "variousartists", "na"}


def credit_keys(d, col):
    keys = set()
    for n in split_names(d.get(col)):
        n = re.sub(r"^(late|pt|pandit|dr|ustad|shri|sri|kavi|sant)\.?\s+", "", n.strip(), flags=re.I)
        k = phonetic(n)
        if k and k not in CREDIT_JUNK:
            keys.add(k)
    return keys


class DSU:
    def __init__(self, n):
        self.p = list(range(n))

    def find(self, i):
        while self.p[i] != i:
            self.p[i] = self.p[self.p[i]]
            i = self.p[i]
        return i

    def union(self, a, b):
        a, b = self.find(a), self.find(b)
        if a != b:
            self.p[max(a, b)] = min(a, b)


def group_by_isrc(rows):
    """Rows sharing any ISRC (cells can hold several) are one recording."""
    dsu = DSU(len(rows))
    first = {}
    for i, d in enumerate(rows):
        for t in isrc_tokens(d.get("ISRC")):
            if t in first:
                dsu.union(first[t], i)
            else:
                first[t] = i
    return _collect(dsu, len(rows))


class Credits:
    """Lyricists and composers of a set of rows."""
    def __init__(self, lyr=(), comp=()):
        self.lyr, self.comp = set(lyr), set(comp)

    def add(self, other):
        self.lyr |= other.lyr
        self.comp |= other.comp

    def conflicts(self, other):
        """Both name lyricists and nobody on one side is credited on the other:
        same title, different song ('Dilbar Dilbar' by Rajendra Krishan vs Sameer)."""
        return bool(self.lyr and other.lyr and not (self.lyr & other.lyr)
                    and not (self.lyr & other.comp) and not (other.lyr & self.comp))


def group_by_song(rows):
    """Recordings of one composition. Joined when they share an ISRC / Spotify /
    Muso track ID, or have the same cleaned title (or a near-identical spelling)
    and their credits do not conflict."""
    n = len(rows)
    dsu = DSU(n)
    seen = {}
    for i, d in enumerate(rows):
        ids = [("isrc", t) for t in isrc_tokens(d.get("ISRC"))]
        for col in ("Spotify Track ID", "Muso Track ID"):
            if not is_placeholder(d.get(col)):
                ids.append((col, str(d[col]).strip()))
        for k in ids:
            if k in seen:
                dsu.union(seen[k], i)
            else:
                seen[k] = i

    # Work on components already joined by ID
    comps = OrderedDict()
    for i in range(n):
        comps.setdefault(dsu.find(i), []).append(i)
    cred = {}
    by_title = defaultdict(list)
    for root, idxs in comps.items():
        c = Credits()
        for i in idxs:
            c.add(Credits(credit_keys(rows[i], "Lyricist"), credit_keys(rows[i], "Composer")))
        cred[root] = c
        for k in OrderedDict.fromkeys(title_key(rows[i].get(TITLE)) for i in idxs):
            if k:
                by_title[k].append(root)

    for roots in by_title.values():
        _cluster_same_title(dsu, roots, cred, comps)

    # Different spellings of one title ("Main Shair To Nahin" / "Main Shayar To Nahin"):
    # needs a shared lyricist, or a shared credit that is not on every song of the report
    freq = Counter(k for c in cred.values() for k in c.lyr | c.comp)
    common = {k for k, v in freq.items() if v > COMMON_CREDIT_SHARE * len(comps)}
    blocks = defaultdict(list)
    for k in by_title:
        blocks[k[:3]].append(k)
    for ks in blocks.values():
        for a in range(len(ks)):
            for b in range(a + 1, len(ks)):
                ka, kb = ks[a], ks[b]
                if min(len(ka), len(kb)) < 8:
                    continue
                sm = difflib.SequenceMatcher(None, ka, kb)
                if sm.quick_ratio() < FUZZY_RATIO or sm.ratio() < FUZZY_RATIO:
                    continue
                for ra in by_title[ka]:
                    for rb in by_title[kb]:
                        x, y = cred[ra], cred[rb]
                        if dsu.find(ra) == dsu.find(rb) or x.conflicts(y):
                            continue
                        if (x.lyr & y.lyr) or ((x.lyr | x.comp) & (y.lyr | y.comp)) - common:
                            dsu.union(ra, rb)
    return _collect(dsu, n)


def _cluster_same_title(dsu, roots, cred, comps):
    """Same title: join components unless their credits conflict. Components with
    no lyricist join the cluster sharing most composers (else the biggest)."""
    clusters = []  # [Credits, size, root]
    with_lyr = sorted((r for r in roots if cred[r].lyr), key=lambda r: -len(comps[r]))
    for r in with_lyr:
        for cl in clusters:
            if not cl[0].conflicts(cred[r]):
                dsu.union(cl[2], r)
                cl[0].add(cred[r]); cl[1] += len(comps[r])
                break
        else:
            c = Credits(); c.add(cred[r])
            clusters.append([c, len(comps[r]), r])
    for r in (r for r in roots if not cred[r].lyr):
        if not clusters:
            c = Credits(); c.add(cred[r])
            clusters.append([c, len(comps[r]), r])
            continue
        best = max(clusters, key=lambda cl: (len(cl[0].comp & cred[r].comp), cl[1]))
        dsu.union(best[2], r)
        best[0].add(cred[r]); best[1] += len(comps[r])


def _collect(dsu, n):
    groups = OrderedDict()
    for i in range(n):
        groups.setdefault(dsu.find(i), []).append(i)
    return list(groups.values())


def primary_order(rows, idxs):
    """Primary version first: untagged title, then earliest release, then most popular."""
    def score(i):
        d = rows[i]
        date = str(d.get("Song Release Date") or "")
        date = date if re.match(r"\d{4}", date) else "9999"
        try:
            pop = -float(d.get("Popularity (0-100)") or 0)
        except (TypeError, ValueError):
            pop = 0
        return (version_type(d.get(TITLE)) != "Original / Re-release", date, pop, i)
    return sorted(idxs, key=score)


def merge_column(h, dicts):
    """Return (original, duplicates) for one column across a group, primary first."""
    if h == "ISRC":
        per_row = [isrc_tokens(d.get(h)) or ([str(d.get(h)).strip()] if not is_placeholder(d.get(h)) else []) for d in dicts]
        key = lambda s: s.upper()
    elif h in LIST_COLS:
        per_row = [split_names(d.get(h)) for d in dicts]
        key = name_key
    else:
        per_row = [[str(d.get(h)).strip()] if not is_placeholder(d.get(h)) else [] for d in dicts]
        key = lambda s: s.casefold()
    seen = set()
    orig, dups = [], []
    for vals in per_row:
        target = orig if not orig else dups
        for v in vals:
            k = key(v)
            if k and k not in seen:
                seen.add(k)
                target.append(v)
    if not orig:
        any_val = next((d.get(h) for d in dicts if d.get(h) not in (None, "")), None)
        return (any_val, None)
    return (", ".join(orig) if h in LIST_COLS or h == "ISRC" else orig[0],
            ", ".join(dups) if dups else None)


def write_merged(wb, title, headers, rows, groups, lead_cols):
    """lead_cols: list of (header, fn(group_no, group_idxs) -> value) shown before the data columns."""
    out_headers, kinds = [], []
    for name, _ in lead_cols:
        out_headers.append(name); kinds.append("isrc" if name in ("Song ID", "ISRC") else "orig")
    for h in headers:
        out_headers.append(h); kinds.append("orig")
        out_headers.append(f"{h} (Duplicates)"); kinds.append("dup")
    ws = wb.create_sheet(title)
    ws.append(out_headers)
    for g, idxs in enumerate(groups):
        dicts = [rows[i] for i in idxs]
        out = [fn(g, idxs) for _, fn in lead_cols]
        for h in headers:
            out.extend(merge_column(h, dicts))
        ws.append(out)
    format_merged(ws, out_headers, kinds)
    return ws


GROUP_FILLS = [PatternFill("solid", fgColor="FFFFFF"), PatternFill("solid", fgColor="DDEBF7")]
ISRC_FILL = PatternFill("solid", fgColor="FFE699")


def write_versions(wb, headers, rows, song_groups, song_ids, vtypes):
    lead = ["Song ID", "Song", "Version Type", "Primary"]
    out_headers = lead + headers
    ws = wb.create_sheet("Song Versions")
    ws.append(out_headers)
    bands = []
    for g, idxs in enumerate(song_groups):
        for n, i in enumerate(idxs):
            d = rows[i]
            ws.append([song_ids[g], rows[idxs[0]].get(TITLE), vtypes[i], "Yes" if n == 0 else ""]
                      + [d.get(h) for h in headers])
            bands.append((g % 2, n == 0))
    format_plain(ws, out_headers)
    isrc_ci = out_headers.index("ISRC") + 1
    for r, (band, primary) in enumerate(bands, start=2):
        for ci in range(1, len(out_headers) + 1):
            c = ws.cell(row=r, column=ci)
            c.fill = ISRC_FILL if ci == isrc_ci else GROUP_FILLS[band]
            if primary:
                c.font = Font(name="Calibri", size=10, bold=True)
    ws.freeze_panes = "C2"


def write_summary(wb, src_rows, isrc_groups, song_groups, rows, song_ids, vtypes):
    ws = wb.create_sheet("Summary", 0)
    multi = [g for g in song_groups if len(g) > 1]
    stats = [
        ("Source rows (Muso Credits)", src_rows),
        ("Unique recordings (by ISRC)", len(isrc_groups)),
        ("Unique songs (by composition)", len(song_groups)),
        ("Songs with more than one version", len(multi)),
        ("Rows folded into another song", src_rows - len(song_groups)),
    ]
    ws.append(["Metric", "Value"])
    for s in stats:
        ws.append(list(s))
    ws.append([])
    ws.append(["Version type", "Rows"])
    for label, cnt in Counter(vtypes).most_common():
        ws.append([label, cnt])
    ws.append([])
    ws.append(["Song ID", "Song (primary version)", "Versions", "Version types"])
    top_hdr = ws.max_row
    order = sorted(range(len(song_groups)), key=lambda g: -len(song_groups[g]))[:25]
    for g in order:
        idxs = song_groups[g]
        types = Counter(vtypes[i] for i in idxs)
        ws.append([song_ids[g], rows[idxs[0]].get(TITLE), len(idxs),
                   ", ".join(f"{t} ×{c}" for t, c in types.most_common())])
    for r in (1, len(stats) + 3, top_hdr):
        for c in ws[r]:
            if c.value is not None:
                c.fill, c.font, c.border = HDR_FILL, HDR_FONT, BORDER
    for col, w in zip("ABCD", (34, 46, 10, 80)):
        ws.column_dimensions[col].width = w


def run_muso_merge(src_path, out_path):
    src_wb = openpyxl.load_workbook(src_path)
    if MUSO_SHEET not in src_wb.sheetnames:
        raise ValueError(f"No '{MUSO_SHEET}' sheet in {src_path}")
    ws = src_wb[MUSO_SHEET]
    headers = [c.value for c in ws[1]]
    rows = [{h: r[i] for i, h in enumerate(headers)}
            for r in ws.iter_rows(min_row=2, values_only=True) if any(r)]

    isrc_groups = group_by_isrc(rows)
    song_groups = [primary_order(rows, g) for g in group_by_song(rows)]
    isrc_groups = [primary_order(rows, g) for g in isrc_groups]
    song_ids = [f"S{n:05d}" for n in range(1, len(song_groups) + 1)]
    vtypes = [version_type(d.get(TITLE)) for d in rows]

    out_wb = openpyxl.Workbook()
    out_wb.remove(out_wb.active)
    write_merged(out_wb, "Merged by Song", headers, rows, song_groups, [
        ("Song ID", lambda g, idxs: song_ids[g]),
        ("Versions", lambda g, idxs: len(idxs)),
        ("Version Types", lambda g, idxs: ", ".join(
            f"{t} ×{c}" for t, c in Counter(vtypes[i] for i in idxs).most_common())),
    ])
    write_versions(out_wb, headers, rows, song_groups, song_ids, vtypes)
    write_merged(out_wb, "Merged by ISRC", headers, rows, isrc_groups, [("Versions", lambda g, idxs: len(idxs))])

    for sh in src_wb.sheetnames:
        s = src_wb[sh]
        sh_headers = [c.value for c in s[1]]
        o = out_wb.create_sheet(sh)
        o.append(sh_headers)
        for r in s.iter_rows(min_row=2, values_only=True):
            if any(r):
                o.append(list(r))
        format_plain(o, sh_headers)

    write_summary(out_wb, len(rows), isrc_groups, song_groups, rows, song_ids, vtypes)
    out_wb.save(out_path)
    return {"out": out_path, "source_rows": len(rows), "recordings": len(isrc_groups),
            "songs": len(song_groups)}


if __name__ == "__main__":
    src = sys.argv[1]
    out = sys.argv[2] if len(sys.argv) > 2 else src.replace(".xlsx", "_merged.xlsx")
    print(run_muso_merge(src, out))
